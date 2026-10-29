from PyPDF2 import PdfReader
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
# 打开 PDF，读取文档结构
reader=PdfReader("./docs/702adb65972a4e28bf00e1057b590c4d.pdf")
parts=[]
# reader.pages 保存各页，逐页提取文字
for page in reader.pages:
    # 某些页面可能没有可提取的文字，用空字符串代替 None
    text=page.extract_text()or ""
    parts.append(text)
text="\n\n".join(parts)
print(text[:500])
print(len(reader.pages))

# 打开已有的 .docx 文档
doc=Document("./docs/纳新题答案与解析.docx")
texts=[]
# # 遍历正文段落
# for paragraph in doc.paragraphs:
#     text=paragraph.text.strip()
#     if text:
#         texts.append(text)
# for table in doc.tables:
#     for row in table.rows:
#         text="\t".join(cell.text for cell in row.cells)
#         texts.append(text)

# content = "\n\n".join(texts)
# print(content[:500])
# 按正文中的先后顺序，取得段落或表格
for block in doc.iter_inner_content():
    if isinstance(block,Paragraph):
        text=block.text.strip()
        if text:
            texts.append(text)
    elif isinstance(block,Table):
        rows=[]
        for row in block.rows:
            row_text="\t".join(cell.text for cell in row.cells)
            if row_text.strip():
                rows.append(row_text)
        text="\n".join(rows)
    if text:
        texts.append(text)
# print(texts[:2])
text="\n\n".join(texts)
# print(text[:16])